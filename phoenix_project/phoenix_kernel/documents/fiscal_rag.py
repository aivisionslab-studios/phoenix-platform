# fiscal_rag.py
import re
from typing import Callable, Optional
from openpyxl import load_workbook
from .confidence import FieldEvidence, ConfidenceStatus

try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
    import numpy as np
    HAS_SKLEARN = True
except ImportError:
    HAS_SKLEARN = False

class FiscalRAG:
    def __init__(self, table_path: str = None):
        self.codes, self.descs, self.tfidf_matrix, self.vectorizer = [], [], None, None
        if table_path: self.load_table(table_path)
            
    def load_table(self, path: str) -> None:
        wb = load_workbook(path, read_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=4, values_only=True):
            if not row or not row[0]: continue
            code = re.sub(r"\D", "", str(row[0]))
            desc = str(row[2] if len(row) > 2 and row[2] else row[1]).strip()
            if len(code) == 8 and desc:
                self.codes.append(code)
                self.descs.append(desc)
        wb.close()
        if HAS_SKLEARN and self.codes:
            self.vectorizer = TfidfVectorizer(ngram_range=(1, 2), min_df=1)
            self.tfidf_matrix = self.vectorizer.fit_transform(self.descs)
            
    def validate_ncm(self, ncm: str) -> bool:
        return ncm in self.codes

    def get_candidates(self, product_name: str, top_k: int = 10) -> list:
        if not HAS_SKLEARN or not self.tfidf_matrix or not product_name: return []
        query_vec = self.vectorizer.transform([product_name])
        similarities = cosine_similarity(query_vec, self.tfidf_matrix).flatten()
        top_indices = np.argsort(similarities)[-top_k:][::-1]
        return [{"code": self.codes[i], "desc": self.descs[i], "score": float(similarities[i])} for i in top_indices if similarities[i] > 0.05]

    def suggest_ncm_with_llm(self, product_name: str, llm_callback: Callable) -> Optional[FieldEvidence]:
        candidates = self.get_candidates(product_name)
        if not candidates: return None
            
        allowlist_str = "\n".join([f"- {c['code']} (Desc: {c['desc']})" for c in candidates])
        code_list = [c["code"] for c in candidates]
        prompt = f"Produto: '{product_name}'.\nOpções NCM:\n{allowlist_str}\nEscolha o NCM MAIS ESPECÍFICO. Evite '.90'. Responda APENAS o código."
        
        llm_response = None
        try: llm_response = llm_callback(prompt)
        except: pass
            
        if llm_response:
            llm_code = re.sub(r"\D", "", str(llm_response))[:8]
            if llm_code in code_list:
                return FieldEvidence(llm_code, ConfidenceStatus.PROBABLE, "llm_tfidf_rag", "Injetado por LLM")
                
        best = candidates[0]
        if not best["code"].endswith("90"):
            return FieldEvidence(best["code"], ConfidenceStatus.PROBABLE, "tfidf_fallback", "Injetado por TF-IDF")
        for c in candidates[1:]:
            if not c["code"].endswith("90"):
                return FieldEvidence(c["code"], ConfidenceStatus.PROBABLE, "tfidf_fallback", "Injetado por TF-IDF (evitando .90)")
        return FieldEvidence(best["code"], ConfidenceStatus.PROBABLE, "tfidf_fallback", "Injetado (genérico)")

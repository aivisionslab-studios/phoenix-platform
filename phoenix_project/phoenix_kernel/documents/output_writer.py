# output_writer.py
import shutil
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from .confidence import is_writeable, needs_audit
from .header_mapper import map_header

_AUDIT_HEADERS = ["Produto", "Campo", "Valor_Encontrado", "Status", "Origem", "Motivo", "Candidatos"]

def write_to_template(template_path: str, output_path: str, records: list, sheet_name: str = None) -> str:
    shutil.copy2(template_path, output_path)
    wb = load_workbook(output_path)
    ws = wb[sheet_name] if sheet_name else wb.active
    col_map = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=c).value
        if h:
            ft, _ = map_header(str(h))
            if ft and ft not in col_map: col_map[ft] = c
            
    start_row = 2
    for r in range(2, ws.max_row + 2):
        if ws.cell(row=r, column=col_map.get("explicit_product_name", 2)).value is None: start_row = r; break
            
    audit_ws = wb["_PHOENIX_AUDIT"] if "_PHOENIX_AUDIT" in wb.sheetnames else wb.create_sheet("_PHOENIX_AUDIT")
    for i, h in enumerate(_AUDIT_HEADERS, 1): audit_ws.cell(row=1, column=i, value=h)
        
    audit_row = 2
    for record in records:
        row_num = start_row; start_row += 1
        for ft, ev in record.items():
            if not hasattr(ev, 'status'): continue
            if is_writeable(ev) and ft in col_map: ws.cell(row=row_num, column=col_map[ft]).value = ev.value
            if needs_audit(ev):
                vals = [record.get("explicit_product_name").value if record.get("explicit_product_name") else "", ft, str(ev.value), ev.status.value, ev.source, ev.reason]
                for i, v in enumerate(vals, 1): audit_ws.cell(row=audit_row, column=i, value=v)
                audit_row += 1
    wb.save(output_path)
    return output_path

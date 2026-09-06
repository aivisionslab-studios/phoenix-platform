import pytest
from openpyxl import Workbook
from phoenix_kernel.documents.pipeline_orchestrator import fill_spreadsheet
from openpyxl import load_workbook

_MIXED_DOC = """
Descrição\tNCM\tCEST\tPeso
Cerveja X\t22030000\t0302100\t0.350
Refrigerante Y\t22021000\t0301000\t0.350

Tinta Acrílica Suvinil Toque de Seda 18L
[DADOS ESTRUTURADOS PARA O ERP] NCM: 32091010 | CEST: 2400100
Peso: 24,5

63. Cachaça São Francisco 970ml
NCM: 22084000
Peso: 0,970
Preço Venda Varejo: R$ 16,90
"""

def test_mixed_document_processing(tmp_path):
    template = tmp_path / "t.xlsx"
    output = tmp_path / "o.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Descrição", "NCM", "CEST", "Peso (Kg)", "Preço Venda Varejo"])
    wb.save(template)
    
    def mock_llm(prompt, allowlist=None): return "22030000"
        
    stats = fill_spreadsheet(_MIXED_DOC, str(template), str(output), "", mock_llm)
    
    assert stats["total"] >= 4
    wb_out = load_workbook(output)
    ws_out = wb_out.active
    
    assert "Cerveja X" in str(ws_out.cell(2, 1).value)
    assert ws_out.cell(2, 2).value == "22030000"
    
    assert "Refrigerante Y" in str(ws_out.cell(3, 1).value)
    
    assert "Suvinil" in str(ws_out.cell(4, 1).value)
    assert ws_out.cell(4, 2).value == "32091010"
    
    assert "Cachaça" in str(ws_out.cell(5, 1).value)
    assert "63." not in str(ws_out.cell(5, 1).value)
    assert ws_out.cell(5, 2).value == "22084000"
    assert ws_out.cell(5, 5).value == 16.90

"""Phoenix Document Pipeline V2 - Fase 2: parser estrutural real de DOCX.

PHX-NEW (2026-08-29, sequência direta da Fase 1 - ver PHX-NEW no topo de
`normalized.py` pro contexto completo): produz um `NormalizedDocument` a
partir de um `.docx`, na ORDEM REAL em que o Word apresenta os elementos
(`document.element.body`, que intercala `w:p` e `w:tbl` na sequência
verdadeira do documento).

Isto é DELIBERADAMENTE um arquivo NOVO e ISOLADO - `_extract_docx()` (em
`documents/engine.py`, usado hoje por todo o fluxo real de leitura de
documento/preenchimento de planilha) continua exatamente como está, sem
nenhuma mudança de comportamento pra quem já usa a Phoenix. `_extract_docx`
tem um bug estrutural conhecido e documentado (despeja TODOS os parágrafos
primeiro e SÓ DEPOIS todas as tabelas, fora de ordem, cada linha de tabela
virando uma string "cel1 | cel2 | cel3" sem cabeçalho) - mas trocar o
resultado dele agora, antes do resto do pipeline (Candidate Engine,
segmentador, executor - Fases 3+) existir pra consumir um `NormalizedDocument`
de verdade, só criaria uma correção isolada que teria que ser refeita
quando essas fases chegarem. Ligar este parser novo ao fluxo real de
produção é trabalho de uma fase posterior, feito com cuidado cirúrgico
(ver instrução permanente do usuário sobre `phoenix_kernel/documents/engine.py`
ser um arquivo com outro trabalho em andamento).

Limitações conhecidas desta primeira versão (documentadas de propósito,
não escondidas):
  - Imagem embutida vira só um `Block` tipo IMAGE (relationship_id +
    metadados de arquivo) - nenhum conteúdo é lido/OCR'ado aqui (isso é
    Fase 12). Quando um parágrafo tem texto E imagem juntos, os blocos de
    imagem daquele parágrafo são emitidos ANTES do bloco de texto - a
    ordem relativa exata entre um "run" de texto e um "run" de imagem
    dentro do MESMO parágrafo não é preservada (só a ordem entre
    parágrafos/tabelas é garantida). Caso raro; refinar se algum
    documento real precisar.
  - Célula de tabela mesclada (merge) é uma particularidade conhecida do
    python-docx (a mesma célula "física" pode aparecer repetida em mais
    de uma posição lógica da linha) - não tratado aqui ainda.
  - A primeira linha de cada tabela é sempre tratada como cabeçalho
    (`headers`); tabelas sem cabeçalho de verdade (ex: uma matriz de
    dados pura) ainda assim têm a primeira linha "roubada" como headers -
    aceitável por ora porque bate com o formato que o resto do desenho
    (Fase 3+) já assume, mas é uma simplificação, não uma detecção real
    de cabeçalho.
  - Parágrafo vazio (sem texto e sem imagem) não vira Block - mesmo
    comportamento de `_extract_docx()` hoje (não carrega informação útil).
    `paragraph_index` ainda assim conta ele (reflete a posição FÍSICA real
    no documento, não só a posição entre os parágrafos que viraram bloco).
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from phoenix_kernel.documents.normalized import (
    Block,
    BlockSource,
    BlockType,
    DocumentMetadata,
    ImageMedia,
    NormalizedDocument,
    compute_document_id,
)

_HEADING_LEVEL_RE = re.compile(r"heading\s*(\d+)", re.IGNORECASE)


def _heading_level_from_style_name(style_name: Optional[str]) -> Optional[int]:
    """Devolve o nível (1, 2, 3...) se `style_name` for um estilo de
    título do Word ("Heading 1", "Heading 2"...), 0 pro estilo "Title",
    ou None se não for um título (parágrafo normal)."""
    if not style_name:
        return None
    match = _HEADING_LEVEL_RE.search(style_name)
    if match:
        return int(match.group(1))
    if style_name.strip().lower() == "title":
        return 0
    return None


def _extract_inline_images(paragraph: Paragraph, document) -> list[tuple[str, object]]:
    """Devolve [(relationship_id, image_part), ...] pra toda imagem
    embutida (inline) dentro deste parágrafo, na ordem em que aparecem no
    XML. Não lê o CONTEÚDO da imagem - só identifica que ela existe e
    onde (relationship id do DOCX + nome/mime type do arquivo embutido)."""
    found: list[tuple[str, object]] = []
    for blip in paragraph._p.findall(".//" + qn("a:blip")):
        embed_id = blip.get(qn("r:embed"))
        if not embed_id:
            continue
        part = document.part.related_parts.get(embed_id)
        if part is not None:
            found.append((embed_id, part))
    return found


def parse_docx_to_normalized_document(path: Path, source_name: Optional[str] = None) -> NormalizedDocument:
    """Lê um `.docx` e devolve um `NormalizedDocument` com um `Block` por
    parágrafo não-vazio/título/tabela/imagem, na ORDEM REAL do documento
    (`document.element.body`, que intercala `w:p` e `w:tbl` na sequência
    verdadeira - ao contrário de `document.paragraphs` + `document.tables`
    separados, que é o que faz `_extract_docx()` hoje e por isso perde a
    ordem relativa entre texto e tabela).

    `document_id` é sempre o hash sha256 do CONTEÚDO do arquivo (nunca o
    caminho - ver `compute_document_id`), então dois arquivos idênticos
    (em pastas ou máquinas diferentes) resolvem pro mesmo id."""
    import docx  # import tardio, mesmo padrão já usado em documents/engine.py

    path = Path(path)
    document = docx.Document(str(path))

    blocks: list[Block] = []
    order = 0
    paragraph_index = 0
    table_index = 0

    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)

            for embed_id, image_part in _extract_inline_images(paragraph, document):
                blocks.append(Block(
                    id=f"b{order:06d}",
                    type=BlockType.IMAGE,
                    order=order,
                    source=BlockSource(paragraph_index=paragraph_index, relationship_id=embed_id),
                    media=ImageMedia(
                        filename=Path(str(image_part.partname)).name,
                        mime_type=image_part.content_type,
                    ),
                ))
                order += 1

            text = paragraph.text
            if text.strip():
                level = _heading_level_from_style_name(paragraph.style.name if paragraph.style is not None else None)
                blocks.append(Block(
                    id=f"b{order:06d}",
                    type=BlockType.HEADING if level is not None else BlockType.PARAGRAPH,
                    order=order,
                    text_raw=text,
                    text_normalized=text,
                    level=level,
                    source=BlockSource(paragraph_index=paragraph_index),
                ))
                order += 1

            paragraph_index += 1

        elif element.tag == qn("w:tbl"):
            table = Table(element, document)
            rows_text = [[cell.text for cell in row.cells] for row in table.rows]
            headers = rows_text[0] if rows_text else []
            data_rows = rows_text[1:] if len(rows_text) > 1 else []
            blocks.append(Block(
                id=f"b{order:06d}",
                type=BlockType.TABLE,
                order=order,
                headers=headers,
                rows=data_rows,
                source=BlockSource(table_index=table_index),
            ))
            order += 1
            table_index += 1

        # Qualquer outro elemento do body (ex: w:sectPr - propriedades de
        # seção/página) é ignorado de propósito - não carrega conteúdo.

    created_at = None
    if document.core_properties.created:
        created_at = document.core_properties.created.isoformat()

    metadata = DocumentMetadata(
        title=document.core_properties.title or None,
        author=document.core_properties.author or None,
        created_at=created_at,
    )

    return NormalizedDocument(
        document_id=compute_document_id(path.read_bytes()),
        source_name=source_name or path.name,
        source_type="docx",
        metadata=metadata,
        blocks=blocks,
    )

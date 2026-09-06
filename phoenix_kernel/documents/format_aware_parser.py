"""Parser ciente de formato — camada de INTEGRAÇÃO (GLM + pipeline testado).

Esta é a ponte entre a detecção de formato genérica (structure_detector +
name_detector, criados na colaboração com o GLM) e o Document Pipeline V2 já
existente e testado (candidate_engine → record_segmenter → identity_engine →
output_writer).

Estratégia (decisão de arquitetura do usuário): a detecção de formato entra
como uma CAMADA NA FRENTE do pipeline, não como substituição. Ela faz o que o
pipeline fazia mal — quebrar o texto em UM bloco por produto, detectando o
formato automaticamente — e entrega um NormalizedDocument (a estrutura que o
pipeline testado já consome). O resto do pipeline (extração de campos, dedupe,
escrita com auditoria, guarda-fiscal) continua idêntico, com seus 612 testes
intactos.

Ganho: antes o `normalized_document_from_text` quebrava por parágrafo (ingênuo)
e marcava título por heurística fraca — o que fazia a coluna Descrição sair
vazia em formatos não previstos e gerava linhas fantasma. Agora a segmentação
é por PRODUTO, com o formato detectado (blocos rotulados, lista numerada,
tabela, título em destaque), e o nome do produto é resolvido pelo detector
genérico. Validado nos dois documentos reais do usuário: conveniência (296
produtos) e construção (125 produtos), com o MESMO código.
"""

from __future__ import annotations

import uuid

from phoenix_kernel.documents.normalized import (
    Block, BlockType, DocumentMetadata, NormalizedDocument,
)
from phoenix_kernel.documents.structure_detector import detect_format, DocumentFormat
from phoenix_kernel.documents.name_detector import detect_name


def format_aware_document_from_text(
    text: str, source_name: str = "", source_type: str = "txt",
) -> tuple[NormalizedDocument, dict]:
    """Constrói um NormalizedDocument segmentando o texto por PRODUTO, usando a
    detecção de formato genérica. Devolve (documento, info_deteccao).

    Cada bloco-de-produto detectado vira UM Block de tipo HEADING contendo o
    texto inteiro do produto, com o nome resolvido no início — o
    record_segmenter e o candidate_engine já sabem processar isso. Blocos sem
    conteúdo real (sem nome e sem dados) são descartados (o filtro de linhas
    fantasma que o pipeline precisava).

    info_deteccao traz: formato detectado, confiança, nº de blocos, nº de
    blocos com nome — útil para log/auditoria e para o usuário entender o que
    o sistema "viu" no documento.
    """
    detection = detect_format(text or "")

    blocks: list[Block] = []
    order = 0
    named = 0
    for det_block in detection.blocks:
        name_ev = detect_name(det_block)
        has_name = name_ev is not None
        if has_name:
            named += 1

        # descarta bloco-fantasma: sem nome E sem dados estruturados não é
        # produto (mesma regra do filtro de fantasmas do pipeline).
        if not has_name and not det_block.has_structured_data:
            continue

        # o texto do bloco começa com o nome resolvido pela detecção genérica,
        # PREFIXADO com o rótulo "Nome do Produto:" — que é como o
        # candidate_engine do pipeline testado reconhece um nome. Sem esse
        # rótulo, o nome que a camada de formato achou (título numerado, caixa
        # alta, destaque) seria jogado fora pelo candidate_engine, que só
        # entende nome por rótulo explícito. Assim a inteligência de detecção
        # de nome do GLM chega intacta ao pipeline testado.
        block_text = det_block.text
        if has_name:
            block_text = f"Nome do Produto: {name_ev.value}\n" + block_text

        blocks.append(Block(
            id=f"fa{order:05d}",
            type=BlockType.HEADING,
            order=order,
            text_raw=block_text,
            text_normalized=block_text,
            level=1,
        ))
        order += 1

    # PHX-NEW (Passo 1 da colaboração "cross-sell/gatilhos"): captura conteúdo
    # de estratégia comercial (matriz de margem por categoria, tipo
    # "Quadrante: O BOI") como dado estruturado, em vez de simplesmente
    # descartar depois que o structure_detector aprendeu a não confundir isso
    # com produto. Roda sobre o TEXTO BRUTO (antes da segmentação em blocos de
    # produto), então não depende de esse conteúdo ter sobrevivido intacto
    # dentro de algum Block. Nunca inventa: documento sem esse padrão -> lista
    # vazia.
    from phoenix_kernel.documents.business_strategy_extractor import extract_pricing_quadrants
    pricing_quadrants = extract_pricing_quadrants(text)

    info = {
        "format": detection.format.value,
        "confidence": round(detection.confidence, 3),
        "blocks_total": len(detection.blocks),
        "blocks_kept": len(blocks),
        "blocks_named": named,
        "pricing_quadrants": pricing_quadrants,
    }

    document = NormalizedDocument(
        document_id=uuid.uuid4().hex,
        source_name=source_name or "documento",
        source_type=source_type,
        metadata=DocumentMetadata(),
        blocks=blocks,
    )
    return document, info


def write_pricing_strategy_sheet(output_xlsx_path: str, quadrants: list) -> bool:
    """Escreve os quadrantes de estratégia capturados como uma aba EXTRA no
    workbook já gerado pelo pipeline — depois que o pipeline principal
    (run_pipeline_docx_to_xlsx) já escreveu a planilha de produtos + a aba
    _PHOENIX_AUDIT. Additive: abre o arquivo, adiciona a aba, salva de volta.
    Nunca mexe nas abas existentes. Devolve False (sem tocar o arquivo) se a
    lista de quadrantes estiver vazia — nada a escrever."""
    if not quadrants:
        return False

    import openpyxl

    wb = openpyxl.load_workbook(output_xlsx_path)
    sheet_name = "_ESTRATEGIA_COMERCIAL"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.append([
        "Quadrante", "Nome", "Subtítulo", "Produtos Exemplo",
        "Margem (rótulo)", "Margem Mín %", "Margem Máx %", "Função/Estratégia",
    ])
    for q in quadrants:
        ws.append([
            q.number, q.name, q.subtitle, "; ".join(q.example_products),
            q.margin_label, q.margin_min_pct, q.margin_max_pct, q.role,
        ])
    wb.save(output_xlsx_path)
    return True


def write_cross_sell_sheet(output_xlsx_path: str, products: list[dict], quadrants: list,
                           *, name_field: str = "Descrição") -> bool:
    """Passo 2: gera sugestões de combo/cross-sell cruzando os produtos reais
    com os quadrantes de estratégia capturados, e escreve como aba EXTRA
    (_SUGESTOES_CROSSSELL) no workbook já gerado. Additive, mesmo padrão de
    write_pricing_strategy_sheet — nunca toca nas outras abas. Sem
    quadrantes ou sem pares âncora/complementar, não escreve nada (devolve
    False) — nunca inventa uma sugestão de negócio."""
    from phoenix_kernel.documents.cross_sell_suggester import suggest_combos

    combos = suggest_combos(products, quadrants, name_field=name_field)
    if not combos:
        return False

    import openpyxl

    wb = openpyxl.load_workbook(output_xlsx_path)
    sheet_name = "_SUGESTOES_CROSSSELL"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.append([
        "Produto Âncora", "Quadrante Âncora", "Produto Complementar",
        "Quadrante Complementar", "Margem do Complementar", "Justificativa",
    ])
    for c in combos:
        ws.append([
            c.anchor_product, c.anchor_quadrant, c.complementary_product,
            c.complementary_quadrant, c.complementary_margin_label, c.justification,
        ])
    wb.save(output_xlsx_path)
    return True

"""Teste de integração: _enhance_sales_triggers_with_llm (Passo 3, opt-in).

Usa um runtime FALSO (não depende de modelo real rodando) para provar que a
integração (resolução de modelo, montagem do prompt, escrita cirúrgica só
nos produtos casados com quadrante, limite de chamadas, fallback em falha)
funciona corretamente.
"""
import asyncio
import tempfile
from pathlib import Path

import openpyxl
import pytest

from core.domain.execution import ExecutionResult, ExecutionStatus
from phoenix_kernel.documents.business_strategy_extractor import PricingQuadrant
from phoenix_kernel.resident.resident_manager import ResidentManager


class _FakeRuntime:
    def __init__(self, output="Frase gerada pelo LLM", should_fail=False):
        self.call_count = 0
        self.output = output
        self.should_fail = should_fail

    async def execute(self, plan):
        self.call_count += 1
        if self.should_fail:
            return ExecutionResult(plan_id="f", status=ExecutionStatus.FAILED, errors=("erro simulado",))
        return ExecutionResult(plan_id="f", status=ExecutionStatus.SUCCESS, output=self.output)


class _FakeRegistry:
    def resolve(self, kind, hint=None):
        class R:
            runtime = "llama.cpp"
            id = "fake-model"
        return R()


def _quadrants():
    return [
        PricingQuadrant(number=1, name="O BOI", subtitle="", example_products=["Cerveja"],
                        margin_label="Margem Mínima", margin_min_pct=5, margin_max_pct=12, role="Isca"),
        PricingQuadrant(number=3, name="A ESTRELA", subtitle="", example_products=["Vinho Casillero"],
                        margin_label="Margem Alta", margin_min_pct=50, margin_max_pct=80, role="Lucro"),
    ]


def _make_xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Descrição", "Descrição do Produto"])
    for r in rows:
        ws.append(r)
    p = Path(tempfile.mkdtemp()) / "t.xlsx"
    wb.save(p)
    return p


def test_only_matched_products_get_rewritten():
    xlsx = _make_xlsx([
        ["Cerveja Heineken 330ml", "<p>determinístico</p>"],
        ["Caderno Espiral", "<p>determinístico</p>"],  # não casa com nenhum quadrante
    ])
    runtime = _FakeRuntime()
    resident = ResidentManager.__new__(ResidentManager)
    resident.runtime = runtime
    resident.registry = _FakeRegistry()

    n = asyncio.run(resident._enhance_sales_triggers_with_llm(str(xlsx), _quadrants(), max_calls=10))
    assert n == 1
    assert runtime.call_count == 1

    wb = openpyxl.load_workbook(xlsx)
    assert "Frase gerada pelo LLM" in wb.active.cell(2, 2).value
    assert wb.active.cell(3, 2).value == "<p>determinístico</p>"  # não mexeu


def test_respects_max_calls_limit():
    xlsx = _make_xlsx([["Cerveja X", ""], ["Vinho Casillero Y", ""], ["Cerveja Z", ""]])
    runtime = _FakeRuntime()
    resident = ResidentManager.__new__(ResidentManager)
    resident.runtime = runtime
    resident.registry = _FakeRegistry()

    n = asyncio.run(resident._enhance_sales_triggers_with_llm(str(xlsx), _quadrants(), max_calls=2))
    assert n == 2
    assert runtime.call_count == 2


def test_llm_failure_keeps_deterministic_text():
    xlsx = _make_xlsx([["Cerveja Heineken 330ml", "<p>texto determinístico mantido</p>"]])
    runtime = _FakeRuntime(should_fail=True)
    resident = ResidentManager.__new__(ResidentManager)
    resident.runtime = runtime
    resident.registry = _FakeRegistry()

    n = asyncio.run(resident._enhance_sales_triggers_with_llm(str(xlsx), _quadrants(), max_calls=10))
    assert n == 0
    wb = openpyxl.load_workbook(xlsx)
    assert wb.active.cell(2, 2).value == "<p>texto determinístico mantido</p>"


def test_no_quadrants_matched_writes_nothing():
    xlsx = _make_xlsx([["Caderno Espiral", "<p>x</p>"]])
    runtime = _FakeRuntime()
    resident = ResidentManager.__new__(ResidentManager)
    resident.runtime = runtime
    resident.registry = _FakeRegistry()

    n = asyncio.run(resident._enhance_sales_triggers_with_llm(str(xlsx), _quadrants(), max_calls=10))
    assert n == 0
    assert runtime.call_count == 0

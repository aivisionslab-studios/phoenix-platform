"""Testes da Fase 2 do "Document Pipeline V2"
(phoenix_kernel/documents/docx_parser.py) - parser estrutural real de
DOCX, que produz um NormalizedDocument (Fase 1) na ORDEM REAL do
documento.

O teste mais importante deste arquivo é
`test_paragraph_table_paragraph_order_is_preserved` - reproduz
exatamente o bug estrutural que motivou toda a investigação: hoje
`_extract_docx()` (documents/engine.py) despeja TODOS os parágrafos e só
DEPOIS todas as tabelas, destruindo a ordem relativa entre texto e
tabela. Aqui provamos que o parser novo NÃO repete esse erro.

Constrói arquivos .docx sintéticos com python-docx em cada teste (não
depende de nenhum arquivo real do usuário) - inclusive uma imagem
embutida via um PNG mínimo (1x1, hardcoded em base64) pra não precisar de
Pillow, que não é dependência da Phoenix."""
from __future__ import annotations

import base64
import io
from pathlib import Path

import docx
import pytest

from phoenix_kernel.documents.docx_parser import parse_docx_to_normalized_document
from phoenix_kernel.documents.normalized import BlockType, NormalizedDocument, compute_document_id

# PNG 1x1 transparente mínimo - python-docx lê dimensão/formato sozinho
# (tem parser de imagem embutido pra PNG/JPEG/GIF/BMP/TIFF), não precisa
# de Pillow instalado.
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _build_docx(path: Path) -> None:
    """Documento sintético cobrindo o caso central da investigação: um
    parágrafo ANTES de uma tabela, e outro DEPOIS dela - se o parser
    reordenar (bug antigo), este teste pega."""
    document = docx.Document()
    document.add_heading("Produtos de construção", level=1)
    document.add_paragraph("paragraph_before_table")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Produto"
    table.cell(0, 1).text = "EAN"
    table.cell(0, 2).text = "Preço"
    table.cell(1, 0).text = "Tinta Suvinil"
    table.cell(1, 1).text = "7891019125301"
    table.cell(1, 2).text = "589,00"
    document.add_paragraph("paragraph_after_table")
    document.add_paragraph("")  # parágrafo vazio - não deve virar Block
    document.add_paragraph("depois da imagem")
    document.save(str(path))


def _build_docx_with_image(path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("antes da imagem")
    document.add_picture(io.BytesIO(_TINY_PNG))
    document.add_paragraph("depois da imagem")
    document.save(str(path))


def _build_docx_with_two_tables(path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("bloco 1")
    t1 = document.add_table(rows=1, cols=2)
    t1.cell(0, 0).text = "A"
    t1.cell(0, 1).text = "B"
    document.add_paragraph("bloco 2")
    t2 = document.add_table(rows=1, cols=2)
    t2.cell(0, 0).text = "C"
    t2.cell(0, 1).text = "D"
    document.add_paragraph("bloco 3")
    document.save(str(path))


@pytest.fixture
def tmp_docx(tmp_path) -> Path:
    path = tmp_path / "teste.docx"
    _build_docx(path)
    return path


def test_paragraph_table_paragraph_order_is_preserved(tmp_docx):
    """O teste central desta investigação inteira: um parágrafo antes de
    uma tabela e outro depois dela devem permanecer NESSA ordem - não
    "todos os parágrafos primeiro, tabela depois" (o bug real de
    `_extract_docx()` hoje)."""
    doc = parse_docx_to_normalized_document(tmp_docx)
    type_and_text = [
        (b.type, b.text_raw if b.type != BlockType.TABLE else "TABLE")
        for b in doc.blocks
    ]
    texts_in_order = [t for _, t in type_and_text]
    assert "paragraph_before_table" in texts_in_order
    assert "paragraph_after_table" in texts_in_order
    before_idx = texts_in_order.index("paragraph_before_table")
    table_idx = texts_in_order.index("TABLE")
    after_idx = texts_in_order.index("paragraph_after_table")
    assert before_idx < table_idx < after_idx


def test_block_order_is_strictly_monotonic(tmp_docx):
    doc = parse_docx_to_normalized_document(tmp_docx)
    orders = [b.order for b in doc.blocks]
    assert orders == sorted(orders)
    assert len(orders) == len(set(orders))  # nenhum order repetido
    assert orders == list(range(len(orders)))  # sequencial a partir de 0


def test_heading_level_is_captured(tmp_docx):
    doc = parse_docx_to_normalized_document(tmp_docx)
    headings = [b for b in doc.blocks if b.type == BlockType.HEADING]
    assert len(headings) == 1
    assert headings[0].text_raw == "Produtos de construção"
    assert headings[0].level == 1


def test_paragraph_keeps_source_paragraph_index(tmp_docx):
    doc = parse_docx_to_normalized_document(tmp_docx)
    paragraphs = [b for b in doc.blocks if b.type in (BlockType.PARAGRAPH, BlockType.HEADING)]
    indices = [b.source.paragraph_index for b in paragraphs]
    assert all(idx is not None for idx in indices)
    assert indices == sorted(indices)  # origem também nunca é reordenada


def test_empty_paragraph_does_not_become_a_block(tmp_docx):
    doc = parse_docx_to_normalized_document(tmp_docx)
    for block in doc.blocks:
        if block.type in (BlockType.PARAGRAPH, BlockType.HEADING):
            assert block.text_raw.strip() != ""


def test_table_headers_and_rows_are_structured_never_flattened(tmp_docx):
    doc = parse_docx_to_normalized_document(tmp_docx)
    tables = [b for b in doc.blocks if b.type == BlockType.TABLE]
    assert len(tables) == 1
    table = tables[0]
    assert table.headers == ["Produto", "EAN", "Preço"]
    assert table.rows == [["Tinta Suvinil", "7891019125301", "589,00"]]
    assert table.source.table_index == 0


def test_two_tables_keep_relative_order_with_surrounding_paragraphs(tmp_path):
    path = tmp_path / "duas_tabelas.docx"
    _build_docx_with_two_tables(path)
    doc = parse_docx_to_normalized_document(path)
    sequence = []
    for b in doc.blocks:
        if b.type == BlockType.TABLE:
            sequence.append(f"table:{b.source.table_index}")
        else:
            sequence.append(b.text_raw)
    assert sequence == ["bloco 1", "table:0", "bloco 2", "table:1", "bloco 3"]


def test_embedded_image_becomes_placeholder_block(tmp_path):
    path = tmp_path / "com_imagem.docx"
    _build_docx_with_image(path)
    doc = parse_docx_to_normalized_document(path)
    images = [b for b in doc.blocks if b.type == BlockType.IMAGE]
    assert len(images) == 1
    image_block = images[0]
    assert image_block.media.mime_type == "image/png"
    assert image_block.source.relationship_id is not None
    assert image_block.text_raw is None  # nenhum OCR nesta fase - só o placeholder


def test_document_id_is_based_on_file_content_not_path(tmp_path):
    path_a = tmp_path / "a.docx"
    path_b = tmp_path / "subpasta" / "b.docx"
    path_b.parent.mkdir()
    _build_docx(path_a)
    _build_docx(path_b)  # mesmo conteúdo, caminho totalmente diferente
    doc_a = parse_docx_to_normalized_document(path_a)
    doc_b = parse_docx_to_normalized_document(path_b)
    assert doc_a.document_id == doc_b.document_id
    assert doc_a.document_id == compute_document_id(path_a.read_bytes())


def test_roundtrip_through_json_preserves_order_and_structure(tmp_docx):
    doc = parse_docx_to_normalized_document(tmp_docx)
    restored = NormalizedDocument.from_json(doc.to_json())
    assert [b.id for b in restored.blocks] == [b.id for b in doc.blocks]
    assert [b.order for b in restored.blocks] == [b.order for b in doc.blocks]
    restored_tables = [b for b in restored.blocks if b.type == BlockType.TABLE]
    original_tables = [b for b in doc.blocks if b.type == BlockType.TABLE]
    assert restored_tables[0].headers == original_tables[0].headers
    assert restored_tables[0].rows == original_tables[0].rows

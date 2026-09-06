"""Prova de ponta a ponta: a guarda de sanidade de negócio protege a escrita
final, através do PIPELINE REAL (não um teste isolado do módulo)."""
import tempfile
from pathlib import Path

import openpyxl

from phoenix_kernel.documents.pipeline_orchestrator import (
    normalized_document_from_text, run_pipeline_docx_to_xlsx,
)


def _template(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Descrição", "Preço de Custo", "Preço Venda Varejo"])
    p = tmp_path / "template.xlsx"
    wb.save(p)
    return p


def test_impossible_price_goes_to_audit_not_cell(tmp_path):
    """Custo (R$ 50) maior que venda (R$ 30) — margem negativa impossível.
    Através do PIPELINE REAL: nem custo nem venda vão pra célula; ambos
    aparecem na aba _PHOENIX_AUDIT."""
    text = (
        "Produto Impossível -- teste\n\n"
        "Preço de Custo: R$ 50,00\n"
        "Preço Venda Varejo: R$ 30,00\n"
    )
    doc = normalized_document_from_text(text, "d", "txt")
    tpl = _template(tmp_path)
    out = tmp_path / "out.xlsx"
    run_pipeline_docx_to_xlsx(document=doc, template_path=str(tpl), output_path=str(out))

    wb = openpyxl.load_workbook(out)
    ws = wb["Sheet1"]
    H = [c.value for c in ws[1]]
    custo_col = H.index("Preço de Custo") + 1
    venda_col = H.index("Preço Venda Varejo") + 1

    # nenhuma célula de preço foi escrita com o valor suspeito
    for r in range(2, ws.max_row + 1):
        assert ws.cell(r, custo_col).value != 50.0
        assert ws.cell(r, venda_col).value != 30.0

    # a auditoria registra os dois, com o motivo
    assert "_PHOENIX_AUDIT" in wb.sheetnames
    audit = wb["_PHOENIX_AUDIT"]
    audit_field_types = [audit.cell(r, 4).value for r in range(2, audit.max_row + 1)]
    assert "cost_price" in audit_field_types
    assert "retail_price" in audit_field_types


def test_coherent_prices_are_written_normally(tmp_path):
    """Preços coerentes (custo < venda) continuam sendo escritos normalmente
    — a guarda não interfere no caso comum."""
    text = (
        "Produto Normal -- teste\n\n"
        "Preço de Custo: R$ 10,00\n"
        "Preço Venda Varejo: R$ 25,00\n"
    )
    doc = normalized_document_from_text(text, "d", "txt")
    tpl = _template(tmp_path)
    out = tmp_path / "out.xlsx"
    result = run_pipeline_docx_to_xlsx(document=doc, template_path=str(tpl), output_path=str(out))
    assert result.rows_written >= 1

    wb = openpyxl.load_workbook(out)
    ws = wb["Sheet1"]
    H = [c.value for c in ws[1]]
    custo_col = H.index("Preço de Custo") + 1
    venda_col = H.index("Preço Venda Varejo") + 1
    valores_custo = [ws.cell(r, custo_col).value for r in range(2, ws.max_row + 1)]
    valores_venda = [ws.cell(r, venda_col).value for r in range(2, ws.max_row + 1)]
    assert 10.0 in valores_custo
    assert 25.0 in valores_venda

"""Testes do preenchimento DETERMINÍSTICO (rota /pipeline-fill).

Regressão do bug real: preencher planilha via LLM (/fill-template) passou de
90min num catálogo real e virou processo órfão. O caminho determinístico faz
o mesmo em segundos, sem LLM — logo não há inferência para travar/orfanar.
"""
import tempfile
from pathlib import Path

import openpyxl

from phoenix_kernel.documents.pipeline_orchestrator import (
    normalized_document_from_text, run_pipeline_docx_to_xlsx,
)
from phoenix_kernel.documents.smart_filler import smart_fill_xlsx


def _template(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Descrição", "NCM", "CEST", "Peso (Kg)", "Tags", "Origem", "Tipo"])
    p = tmp_path / "template.xlsx"
    wb.save(p)
    return p


def test_deterministic_fill_uses_no_llm_and_is_fast(tmp_path):
    """O caminho determinístico não chama nenhum modelo — é só parser + regras.
    Prova que preenche a partir do texto sem inferência."""
    text = (
        "Tinta Acrílica Suvinil 18L -- Premium\n\n"
        "NCM: 32091010 | CEST: 2400100 | Peso (Kg): 24,5\n\n"
        "Argamassa AC3 Quartzolit 20kg -- Flexível\n\n"
        "NCM: 38245000 | CEST: 1000100 | Peso (Kg): 20,0\n\n"
    )
    doc = normalized_document_from_text(text, "catalogo", "txt")
    tpl = _template(tmp_path)
    out = tmp_path / "out.xlsx"

    import time
    t0 = time.perf_counter()
    result = run_pipeline_docx_to_xlsx(document=doc, template_path=str(tpl), output_path=str(out))
    elapsed = time.perf_counter() - t0

    assert result.rows_written >= 1
    assert elapsed < 10  # determinístico: segundos, não minutos
    assert out.exists()


def test_deterministic_fill_then_smart_fill_completes_empty_cells(tmp_path):
    text = "Cerveja Heineken 330ml -- Lager\n\nNCM: 22030000 | Peso (Kg): 0,33\n\n"
    doc = normalized_document_from_text(text, "cat", "txt")
    tpl = _template(tmp_path)
    out = tmp_path / "out.xlsx"

    run_pipeline_docx_to_xlsx(document=doc, template_path=str(tpl), output_path=str(out))
    report = smart_fill_xlsx(str(out), str(out))

    # smart filler preencheu Origem/Tipo por regra
    wb = openpyxl.load_workbook(out)
    ws = wb["Sheet1"]
    headers = [c.value for c in ws[1]]
    origem_col = headers.index("Origem") + 1
    # ao menos uma linha ganhou Origem=NACIONAL por derivação
    origens = [ws.cell(r, origem_col).value for r in range(2, ws.max_row + 1)]
    assert "NACIONAL" in origens
    assert sum(report.filled_by_rule.values()) >= 1

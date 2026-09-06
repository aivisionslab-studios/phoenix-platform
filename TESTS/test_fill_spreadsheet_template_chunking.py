# TESTS/test_fill_spreadsheet_template_chunking.py
#
# PHX-NEW (achado real do usuário 2026-08-28, testando o próprio recurso
# na sua máquina com um documento de verdade: um catálogo de construção
# extraído de uma conversa longa com o Gemini, ~58KB de texto depois de
# extraído). fill_spreadsheet_template_direct cortava o documento-fonte
# em 18000 caracteres ANTES de mandar pro modelo - o resto era descartado
# em SILÊNCIO. Resultado real: de ~20 produtos distintos mencionados no
# documento, só 3 viraram linha na planilha - um "sucesso" técnico (sem
# erro nenhum na resposta) que escondia uma perda grande de dados. O
# usuário só percebeu abrindo o arquivo final e contando manualmente as
# linhas - a resposta da API não dava nenhum sinal de que faltava algo.
#
# Este arquivo testa a correção: _split_text_into_chunks() (documento
# completo dividido em pedaços que cabem no contexto do modelo, sem
# cortar parágrafo ao meio), _row_identity_key()/_merge_and_dedupe_rows()
# (uma conversa longa costuma repetir o mesmo produto várias vezes ao
# longo de revisões - cada repetição pode cair num pedaço diferente, e a
# versão mais tardia deve vencer), e o fluxo ponta a ponta de
# fill_spreadsheet_template_direct processando um documento GRANDE
# (maior que um pedaço só) com múltiplas chamadas ao runtime mockado.
#
# Rodar com: pytest -q (de dentro de 'PHOENIX 3.0/')

import asyncio
import sys
from pathlib import Path
from unittest.mock import AsyncMock

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

openpyxl = pytest.importorskip("openpyxl")

from core.domain.execution import ExecutionResult, ExecutionStatus  # noqa: E402
from phoenix_kernel.resident.resident_manager import (  # noqa: E402
    ResidentManager,
    DOCUMENT_CREATE_TIMEOUT_LARGE_SECONDS,
    DOCUMENT_CREATE_TIMEOUT_MEDIUM_SECONDS,
    _SPREADSHEET_FILL_CHUNK_CHAR_BUDGET,
    _merge_and_dedupe_rows,
    _row_identity_key,
    _split_text_into_chunks,
)
from phoenix_kernel.models.registry import ModelRegistry  # noqa: E402


# ============================================================
# _split_text_into_chunks
# ============================================================

def test_split_empty_text_returns_no_chunks():
    assert _split_text_into_chunks("", 1000) == []
    assert _split_text_into_chunks(None, 1000) == []  # type: ignore[arg-type]


def test_split_text_smaller_than_budget_is_a_single_chunk():
    text = "Um documento pequeno, cabe inteiro num pedaço só."
    assert _split_text_into_chunks(text, 1000) == [text]


def test_split_keeps_small_paragraphs_together_in_one_chunk():
    text = "Parágrafo um.\n\nParágrafo dois.\n\nParágrafo três."
    chunks = _split_text_into_chunks(text, 1000)
    assert chunks == [text]  # tudo cabe junto, não precisa dividir


def test_split_never_cuts_a_paragraph_in_half():
    para_a = "A" * 15000
    para_b = "B" * 15000
    text = f"{para_a}\n\n{para_b}"
    chunks = _split_text_into_chunks(text, 22000)
    assert len(chunks) == 2
    assert chunks[0] == para_a
    assert chunks[1] == para_b
    # Confirma que nenhum pedaço contém uma MISTURA dos dois parágrafos
    # (prova real de que o corte respeitou a fronteira do parágrafo).
    assert "B" not in chunks[0]
    assert "A" not in chunks[1]


def test_split_hard_slices_a_single_paragraph_bigger_than_the_budget():
    huge_paragraph = "X" * 50000  # um único "parágrafo" sem nenhuma quebra
    chunks = _split_text_into_chunks(huge_paragraph, 22000)
    assert len(chunks) == 3  # 22000 + 22000 + 6000
    assert all(len(c) <= 22000 for c in chunks)
    assert "".join(chunks) == huge_paragraph  # nada foi perdido no processo


def test_split_falls_back_to_single_newline_when_no_blank_lines():
    line_a = "A" * 15000
    line_b = "B" * 15000
    text = f"{line_a}\n{line_b}"  # só uma quebra de linha simples, sem parágrafo com linha em branco
    chunks = _split_text_into_chunks(text, 22000)
    assert len(chunks) == 2
    assert chunks[0] == line_a
    assert chunks[1] == line_b


# ============================================================
# _row_identity_key / _merge_and_dedupe_rows
# ============================================================

def test_row_identity_prefers_barcode_like_value():
    row = {"Descrição": "Produto X", "Código de Barras": "7891019125301"}
    assert _row_identity_key(row) == "ean:7891019125301"


def test_row_identity_falls_back_to_name_like_column():
    row = {"Nome": "Tinta Suvinil", "Preço": "R$ 10,00"}
    assert _row_identity_key(row) == "name:nome:tinta suvinil"


def test_row_identity_is_unique_without_barcode_or_name_column():
    row_a = {"Col1": "valor", "Col2": "outro"}
    row_b = {"Col1": "valor", "Col2": "outro"}
    # Mesmo com valores idênticos, sem um sinal confiável de identidade
    # (nem código de barras, nem coluna de nome/descrição) cada linha é
    # tratada como única - nunca arrisca fundir por engano.
    assert _row_identity_key(row_a) != _row_identity_key(row_b)


def test_merge_keeps_distinct_products_from_different_chunks():
    chunk1 = [{"Nome": "Tinta Suvinil", "Preço": "10"}]
    chunk2 = [{"Nome": "Argamassa Quartzolit", "Preço": "20"}]
    merged = _merge_and_dedupe_rows([chunk1, chunk2])
    assert merged == chunk1 + chunk2


def test_merge_dedupes_same_product_keeping_the_later_chunk_version():
    # Reproduz o achado real: o mesmo produto é mencionado num pedaço
    # ANTERIOR do documento (rascunho) e revisado num pedaço POSTERIOR
    # (versão final, com dados mais completos) - a versão final deve
    # vencer.
    chunk1 = [{"Nome": "Tinta Suvinil", "Preço": "10"}]
    chunk2 = [{"Nome": "Tinta Suvinil", "Preço": "12"}]
    merged = _merge_and_dedupe_rows([chunk1, chunk2])
    assert merged == [{"Nome": "Tinta Suvinil", "Preço": "12"}]


def test_merge_dedupes_by_barcode_across_chunks_even_with_different_names():
    # O nome mudou entre revisões (comum numa conversa de marketing que
    # ajusta o título do produto), mas o código de barras é o mesmo - a
    # Phoenix não pode duplicar a linha só porque o texto do nome mudou.
    chunk1 = [{"Nome": "Tinta Suvinil 18L", "Código de Barras": "7891260021347"}]
    chunk2 = [{"Nome": "Tinta Suvinil Toque de Seda 18L - Edição Final", "Código de Barras": "7891260021347"}]
    merged = _merge_and_dedupe_rows([chunk1, chunk2])
    assert len(merged) == 1
    assert merged[0]["Nome"] == "Tinta Suvinil Toque de Seda 18L - Edição Final"


# ============================================================
# fill_spreadsheet_template_direct - documento grande (múltiplos pedaços)
# ============================================================

def _make_resident() -> ResidentManager:
    resident = object.__new__(ResidentManager)
    resident.registry = ModelRegistry()
    resident.runtime = AsyncMock()
    resident.logs = AsyncMock()
    resident.logs.add_event = lambda *a, **k: None
    resident.ahde = None
    resident._active_models = {}
    return resident


def _make_template_xlsx(path: Path, headers=("Nome", "Código de Barras", "Preço")) -> None:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Produtos"
    for col_idx, h in enumerate(headers, start=1):
        sheet.cell(row=1, column=col_idx, value=h)
    wb.save(str(path))


def _json_output(sheet: str, rows: list[dict]) -> ExecutionResult:
    import json
    return ExecutionResult(plan_id="p1", status=ExecutionStatus.SUCCESS, output=json.dumps({"sheet": sheet, "rows": rows}))


def _make_large_source_docx(path: Path) -> None:
    """Gera um DOCX com dois parágrafos GRANDES - cada um sozinho abaixo
    do orçamento por pedaço (não sofre corte interno), mas OS DOIS JUNTOS
    ultrapassam o orçamento - depois de extraído (parágrafos juntados com
    '\\n\\n', ver _extract_docx), isso obrigatoriamente vira exatamente 2
    pedaços em _SPREADSHEET_FILL_CHUNK_CHAR_BUDGET."""
    import docx
    doc = docx.Document()
    target_len = int(_SPREADSHEET_FILL_CHUNK_CHAR_BUDGET * 0.6)  # cada parágrafo sozinho cabe folgado
    filler = ("lorem ipsum dolor sit amet consectetur adipiscing elit sed do eiusmod " * 500)
    para1 = ("PRODUTO 1 (rascunho inicial): Tinta Suvinil 18L, preço R$ 10,00. " + filler)[:target_len]
    para2 = ("PRODUTO 2 (item novo, só aparece na parte final): Argamassa Quartzolit 20kg. " + filler)[:target_len]
    doc.add_paragraph(para1)
    doc.add_paragraph(para2)
    doc.save(str(path))


def test_large_document_is_split_and_all_chunks_are_queried(tmp_path):
    resident = _make_resident()
    source = tmp_path / "conversa_longa.docx"
    template = tmp_path / "template.xlsx"
    _make_large_source_docx(source)
    _make_template_xlsx(template)

    resident.runtime.execute = AsyncMock(side_effect=[
        _json_output("Produtos", [{"Nome": "Tinta Suvinil 18L", "Código de Barras": "", "Preço": "10"}]),
        _json_output("Produtos", [{"Nome": "Argamassa Quartzolit 20kg", "Código de Barras": "", "Preço": "20"}]),
    ])

    result = asyncio.run(resident.fill_spreadsheet_template_direct(str(source), str(template)))

    assert result["ok"] is True, result
    # O ponto central do achado: as DUAS chamadas (uma por pedaço)
    # aconteceram, e as linhas das duas foram pro arquivo final -
    # nenhuma delas ficou de fora por causa do corte antigo em 18000
    # caracteres.
    assert resident.runtime.execute.call_count == 2
    assert result["chunks_total"] == 2
    assert result["chunks_processed"] == 2
    assert result["rows_written"] == 2

    out_path = Path(result["file_path"])
    wb = openpyxl.load_workbook(str(out_path))
    sheet = wb["Produtos"]
    written_names = {sheet.cell(row=r, column=1).value for r in (2, 3)}
    assert written_names == {"Tinta Suvinil 18L", "Argamassa Quartzolit 20kg"}


def test_large_document_dedupes_product_revised_in_a_later_chunk(tmp_path):
    resident = _make_resident()
    source = tmp_path / "conversa_longa.docx"
    template = tmp_path / "template.xlsx"
    _make_large_source_docx(source)
    _make_template_xlsx(template)

    # Os DOIS pedaços mencionam o "mesmo" produto (mesmo código de
    # barras) - a versão do segundo pedaço (mais tardia) tem que vencer,
    # sem duplicar a linha.
    resident.runtime.execute = AsyncMock(side_effect=[
        _json_output("Produtos", [{"Nome": "Tinta Suvinil (rascunho)", "Código de Barras": "7891260021347", "Preço": "10"}]),
        _json_output("Produtos", [{"Nome": "Tinta Suvinil Toque de Seda 18L", "Código de Barras": "7891260021347", "Preço": "12"}]),
    ])

    result = asyncio.run(resident.fill_spreadsheet_template_direct(str(source), str(template)))

    assert result["ok"] is True, result
    assert result["rows_written"] == 1  # deduplicado, não 2

    out_path = Path(result["file_path"])
    wb = openpyxl.load_workbook(str(out_path))
    sheet = wb["Produtos"]
    assert sheet.cell(row=2, column=1).value == "Tinta Suvinil Toque de Seda 18L"
    assert sheet.cell(row=2, column=3).value == "12"


def test_one_failed_chunk_does_not_abort_the_whole_document(tmp_path):
    resident = _make_resident()
    source = tmp_path / "conversa_longa.docx"
    template = tmp_path / "template.xlsx"
    _make_large_source_docx(source)
    _make_template_xlsx(template)

    # Primeiro pedaço falha (JSON inválido em ambas as tentativas),
    # segundo pedaço funciona - o documento inteiro não pode falhar só
    # porque UM pedaço teve problema; o que der certo deve ser salvo.
    resident.runtime.execute = AsyncMock(side_effect=[
        ExecutionResult(plan_id="p1", status=ExecutionStatus.SUCCESS, output="isto não é JSON"),
        ExecutionResult(plan_id="p1", status=ExecutionStatus.SUCCESS, output="isto não é JSON"),
        _json_output("Produtos", [{"Nome": "Argamassa Quartzolit 20kg", "Código de Barras": "", "Preço": "20"}]),
    ])

    result = asyncio.run(resident.fill_spreadsheet_template_direct(str(source), str(template)))

    assert result["ok"] is True, result
    assert result["rows_written"] == 1
    assert result["chunks_failed"] == 1
    assert len(result["chunks_failed_errors"]) == 1


def test_max_chunks_safety_cap_is_reported_not_silently_dropped(tmp_path, monkeypatch):
    # Documento absurdamente grande (upload por engano, por exemplo) -
    # o teto de segurança existe pra não gerar centenas de chamadas ao
    # modelo, mas a parte que ficou de fora precisa aparecer na resposta
    # (`source_truncated_extra_parts`) - nunca sumir em silêncio como o
    # corte antigo em 18000 caracteres fazia.
    monkeypatch.setattr(
        "phoenix_kernel.resident.resident_manager._SPREADSHEET_FILL_MAX_CHUNKS", 1,
    )
    resident = _make_resident()
    source = tmp_path / "conversa_longa.docx"
    template = tmp_path / "template.xlsx"
    _make_large_source_docx(source)  # gera 2 pedaços reais
    _make_template_xlsx(template)

    resident.runtime.execute = AsyncMock(return_value=_json_output(
        "Produtos", [{"Nome": "Tinta Suvinil 18L", "Código de Barras": "", "Preço": "10"}],
    ))

    result = asyncio.run(resident.fill_spreadsheet_template_direct(str(source), str(template)))

    assert result["ok"] is True, result
    assert resident.runtime.execute.call_count == 1  # só o pedaço permitido pelo teto
    assert result["chunks_total"] == 1
    assert result["source_truncated_extra_parts"] == 1


def test_small_document_still_uses_exactly_one_chunk_and_two_calls_on_bad_json(tmp_path):
    # Controle de regressão: um documento pequeno (o caso comum, e o
    # único testado antes desta correção) precisa continuar se
    # comportando EXATAMENTE como antes - um pedaço só, retry de 1 vez em
    # JSON inválido, sem nenhuma chamada extra.
    resident = _make_resident()
    source = tmp_path / "pequeno.docx"
    template = tmp_path / "template.xlsx"
    import docx
    docx.Document().paragraphs  # no-op, só documentando a intenção
    doc = docx.Document()
    doc.add_paragraph("Cliente: Ana Silva, valor R$ 100,00.")
    doc.save(str(source))
    _make_template_xlsx(template, headers=("Nome", "Valor"))

    resident.runtime.execute = AsyncMock(return_value=ExecutionResult(
        plan_id="p1", status=ExecutionStatus.SUCCESS, output="isto não é JSON de jeito nenhum",
    ))

    result = asyncio.run(resident.fill_spreadsheet_template_direct(str(source), str(template)))

    assert result["ok"] is False
    assert "JSON" in result["error"]
    assert resident.runtime.execute.call_count == 2


# ============================================================
# Orçamento de timeout por pedaço escala com o tamanho do modelo
# (achado real do usuário 2026-08-28, verificando os "30 minutos"
# espalhados pelo código depois do fix de chunking - ver comentário
# completo em fill_spreadsheet_template_direct, onde
# _chunk_document_timeout é calculado).
#
# create_document_direct (função irmã) já escolhia entre
# DOCUMENT_CREATE_TIMEOUT_LARGE_SECONDS (30min, modelo 12b+) e
# DOCUMENT_CREATE_TIMEOUT_MEDIUM_SECONDS (20min, demais modelos) - o
# driver llama_cpp.py também calibra seu próprio piso HTTP por esse
# mesmo critério ("Resident corta em 30 min" no comentário do driver).
# fill_spreadsheet_template_direct nunca fazia essa escolha - usava
# sempre o teto MEDIUM (20min) por pedaço, não importa o modelo. Com o
# modelo padrão (qwen3-8b) isso nunca deu problema visível (o driver
# também cai no piso de 20min pra esse caso - mesmo valor dos dois
# lados). Mas um modelo 12b+ faria o driver calibrar um piso de 1740s
# (pensado pra um orçamento de 30min) enquanto o loop de pedaços
# cancelaria a chamada bem antes, em 1200s - a chamada seria cortada
# ANTES do próprio piso que o driver escolheu pra ela. Estes dois
# testes prova que a nova lógica evita esse descompasso, sem mudar o
# caso comum (modelo pequeno/médio).
# ============================================================

def test_default_small_model_keeps_the_original_medium_timeout_budget(tmp_path):
    # Controle de regressão: SEM um modelo "grande" resolvido, o
    # orçamento por pedaço continua sendo exatamente o de antes desta
    # correção (MEDIUM_SECONDS - 60 = 1140s) - nenhuma mudança de
    # comportamento pro caso comum (qwen3-8b, o modelo real em uso).
    resident = _make_resident()
    source = tmp_path / "pequeno.docx"
    template = tmp_path / "template.xlsx"
    import docx
    doc = docx.Document()
    doc.add_paragraph("Cliente: Ana Silva, valor R$ 100,00.")
    doc.save(str(source))
    _make_template_xlsx(template, headers=("Nome", "Valor"))

    captured_plans = []

    async def _capture(plan):
        captured_plans.append(plan)
        return _json_output("Nome", [{"Nome": "Ana Silva", "Valor": "100"}])

    resident.runtime.execute = AsyncMock(side_effect=_capture)

    result = asyncio.run(resident.fill_spreadsheet_template_direct(str(source), str(template)))

    assert result["ok"] is True, result
    assert len(captured_plans) == 1
    assert captured_plans[0].parameters["timeout_seconds"] == DOCUMENT_CREATE_TIMEOUT_MEDIUM_SECONDS - 60


def test_large_model_gets_the_larger_timeout_budget_per_chunk(tmp_path, monkeypatch):
    # Achado real: forçando a resolução de um modelo "12b+" (mesma lista
    # de marcadores que create_document_direct e llama_cpp.py já usam),
    # cada pedaço agora recebe o orçamento LARGE (30min), consistente com
    # o piso de 1740s que o driver calibraria pra esse mesmo modelo -
    # antes desta correção, isso ficava preso no orçamento MEDIUM (20min)
    # mesmo pra um modelo 12b+.
    resident = _make_resident()
    from types import SimpleNamespace
    resident.registry.resolve = lambda role, hint="": SimpleNamespace(
        id="qwen2.5-14b-instruct", runtime="llama.cpp", roles=(role,),
    )

    source = tmp_path / "pequeno.docx"
    template = tmp_path / "template.xlsx"
    import docx
    doc = docx.Document()
    doc.add_paragraph("Cliente: Ana Silva, valor R$ 100,00.")
    doc.save(str(source))
    _make_template_xlsx(template, headers=("Nome", "Valor"))

    captured_plans = []

    async def _capture(plan):
        captured_plans.append(plan)
        return _json_output("Nome", [{"Nome": "Ana Silva", "Valor": "100"}])

    resident.runtime.execute = AsyncMock(side_effect=_capture)

    result = asyncio.run(resident.fill_spreadsheet_template_direct(str(source), str(template)))

    assert result["ok"] is True, result
    assert len(captured_plans) == 1
    assert captured_plans[0].model == "qwen2.5-14b-instruct"
    assert captured_plans[0].parameters["timeout_seconds"] == DOCUMENT_CREATE_TIMEOUT_LARGE_SECONDS - 60

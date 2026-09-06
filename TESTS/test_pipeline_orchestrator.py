"""Testes do orquestrador do Document Pipeline V2 (o fio que liga as peças)."""

import tempfile
from pathlib import Path

import openpyxl
import pytest

from phoenix_kernel.documents.pipeline_orchestrator import (
    normalized_document_from_text,
    resolve_columns_from_template,
    run_pipeline_docx_to_xlsx,
)


def _marketup_like_template(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Descrição", "NCM", "CEST", "Peso (Kg)", "Tags", "Preço De"])
    p = tmp_path / "template.xlsx"
    wb.save(p)
    return p


def test_columns_resolved_automatically_from_headers(tmp_path):
    tpl = _marketup_like_template(tmp_path)
    columns, header_row, sheet, mapped, unmapped = resolve_columns_from_template(str(tpl))
    field_types = {c.field_type for c in columns}
    assert "ncm" in field_types
    assert "cest" in field_types
    assert "weight" in field_types
    assert "tags" in field_types
    assert header_row == 1
    # PHX-FIX (2026-09-03): "Descrição" agora É mapeada — recebe o nome do
    # produto (explicit_product_name). Antes ficava vazia, o que fazia a
    # planilha sair sem os nomes dos produtos.
    assert "explicit_product_name" in field_types


def test_pipeline_writes_deterministically_with_audit(tmp_path):
    tpl = _marketup_like_template(tmp_path)
    # documento com 2 produtos e dados estruturados
    text = (
        "Tinta Suvinil Toque de Seda 18L -- Acetinado\n\n"
        "NCM: 32091010 | CEST: 2400100 | Peso (Kg): 24,500\n\n"
        "Argamassa AC3 Quartzolit 20kg -- Flexível\n\n"
        "NCM: 38245000 | CEST: 1000100 | Peso (Kg): 20,000\n\n"
    )
    doc = normalized_document_from_text(text, "catalogo", "txt")
    out = tmp_path / "out.xlsx"
    res = run_pipeline_docx_to_xlsx(document=doc, template_path=str(tpl), output_path=str(out))

    assert res.rows_written >= 1
    assert out.exists()
    wb = openpyxl.load_workbook(out)
    # aba de auditoria sempre presente (mesmo que vazia de problemas)
    assert "_PHOENIX_AUDIT" in wb.sheetnames
    # NCM colado na coluna certa
    ws = wb["Sheet1"]
    headers = [c.value for c in ws[1]]
    ncm_col = headers.index("NCM") + 1
    ncm_values = [ws.cell(r, ncm_col).value for r in range(2, ws.max_row + 1)]
    ncm_str = [str(v) for v in ncm_values if v]
    assert any("32091010" in v for v in ncm_str)


def test_pipeline_is_deterministic(tmp_path):
    """Mesma entrada -> mesma saída (é a garantia do caminho sem LLM)."""
    tpl = _marketup_like_template(tmp_path)
    text = "Produto X -- teste\n\nNCM: 12345678 | Peso (Kg): 1,500\n\n"
    doc = normalized_document_from_text(text, "d", "txt")

    out1 = tmp_path / "o1.xlsx"
    out2 = tmp_path / "o2.xlsx"
    r1 = run_pipeline_docx_to_xlsx(document=normalized_document_from_text(text, "d", "txt"),
                                   template_path=str(tpl), output_path=str(out1))
    r2 = run_pipeline_docx_to_xlsx(document=normalized_document_from_text(text, "d", "txt"),
                                   template_path=str(tpl), output_path=str(out2))
    assert r1.rows_written == r2.rows_written
    assert r1.records_total == r2.records_total


def test_template_never_overwritten(tmp_path):
    tpl = _marketup_like_template(tmp_path)
    import hashlib
    before = hashlib.sha256(tpl.read_bytes()).hexdigest()
    doc = normalized_document_from_text("P -- t\n\nNCM: 11112222\n\n", "d", "txt")
    run_pipeline_docx_to_xlsx(document=doc, template_path=str(tpl), output_path=str(tmp_path / "o.xlsx"))
    after = hashlib.sha256(tpl.read_bytes()).hexdigest()
    assert before == after  # template intocado

# TESTS/test_fill_spreadsheet_template_bridge.py
#
# PHX-NEW (pedido do usuário 2026-08-28: "aceitar dois arquivos - um fonte,
# um template alvo - e editar o segundo em vez de gerar do zero, pra
# qualquer documento que a Phoenix já lê e edita").
#
# Testa ResidentManager.fill_spreadsheet_template_direct() de ponta a
# ponta: extração real do documento-fonte (DOCX), leitura real da
# estrutura do template (.xlsx via openpyxl) e preenchimento real do
# template, com o runtime.execute() (LLM) mockado devolvendo um JSON -
# mesmo padrão de test_document_engine_bridge.py
# (_make_resident_for_document_tests: registry real, runtime mockado).
#
# Rodar com: pytest -q (de dentro de 'PHOENIX 3.0/')

import asyncio
import sys
from pathlib import Path
from unittest.mock import AsyncMock

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

openpyxl = pytest.importorskip("openpyxl")

from core.domain.execution import ExecutionResult, ExecutionStatus  # noqa: E402
from phoenix_kernel.resident.resident_manager import ResidentManager  # noqa: E402
from phoenix_kernel.models.registry import ModelRegistry  # noqa: E402


def _make_resident() -> ResidentManager:
    resident = object.__new__(ResidentManager)
    resident.registry = ModelRegistry()
    resident.runtime = AsyncMock()
    resident.logs = AsyncMock()
    resident.logs.add_event = lambda *a, **k: None
    resident.ahde = None
    resident._active_models = {}
    return resident


def _make_source_docx(path: Path) -> None:
    import docx
    doc = docx.Document()
    doc.add_paragraph("Cliente: Ana Silva, CPF 111.111.111-11, valor devido R$ 100,00.")
    doc.add_paragraph("Cliente: Beto Souza, CPF 222.222.222-22, valor devido R$ 200,00.")
    doc.save(str(path))


def _make_template_xlsx(path: Path, headers=("Nome", "CPF", "Valor")) -> None:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Cobranças"
    for col_idx, h in enumerate(headers, start=1):
        sheet.cell(row=1, column=col_idx, value=h)
    wb.save(str(path))


def _json_plan_response(sheet: str, rows: list[dict]) -> ExecutionResult:
    import json
    return ExecutionResult(
        plan_id="p1", status=ExecutionStatus.SUCCESS,
        output=json.dumps({"sheet": sheet, "rows": rows}),
    )


def test_fills_template_from_source_document(tmp_path):
    resident = _make_resident()
    source = tmp_path / "cobrancas.docx"
    template = tmp_path / "modelo.xlsx"
    _make_source_docx(source)
    _make_template_xlsx(template)

    resident.runtime.execute = AsyncMock(return_value=_json_plan_response(
        "Cobranças",
        [
            {"Nome": "Ana Silva", "CPF": "111.111.111-11", "Valor": "R$ 100,00"},
            {"Nome": "Beto Souza", "CPF": "222.222.222-22", "Valor": "R$ 200,00"},
        ],
    ))

    result = asyncio.run(resident.fill_spreadsheet_template_direct(
        str(source), str(template), instruction="Preencha a planilha com os clientes do documento.",
    ))

    assert result["ok"] is True
    assert result["rows_written"] == 2
    assert result["sheet_used"] == "Cobranças"
    assert result["auto_created_columns"] == {}
    assert result["source_file"] == "cobrancas.docx"
    assert result["template_file"] == "modelo.xlsx"

    out_path = Path(result["file_path"])
    assert out_path.exists()
    wb = openpyxl.load_workbook(str(out_path))
    sheet = wb["Cobranças"]
    assert [c.value for c in sheet[2]] == ["Ana Silva", "111.111.111-11", "R$ 100,00"]
    assert [c.value for c in sheet[3]] == ["Beto Souza", "222.222.222-22", "R$ 200,00"]
    # Template original nunca é sobrescrito.
    assert openpyxl.load_workbook(str(template))["Cobranças"].max_row == 1


def test_template_must_be_xlsx(tmp_path):
    resident = _make_resident()
    source = tmp_path / "cobrancas.docx"
    template = tmp_path / "modelo.docx"
    _make_source_docx(source)
    import docx
    docx.Document().save(str(template))

    result = asyncio.run(resident.fill_spreadsheet_template_direct(str(source), str(template)))

    assert result["ok"] is False
    assert ".xlsx" in result["error"]
    resident.runtime.execute.assert_not_called()


def test_missing_source_file_fails_without_calling_runtime(tmp_path):
    resident = _make_resident()
    template = tmp_path / "modelo.xlsx"
    _make_template_xlsx(template)

    result = asyncio.run(resident.fill_spreadsheet_template_direct(
        str(tmp_path / "nao-existe.docx"), str(template),
    ))

    assert result["ok"] is False
    assert "não encontrado" in result["error"]
    resident.runtime.execute.assert_not_called()


def test_missing_template_file_fails_without_calling_runtime(tmp_path):
    resident = _make_resident()
    source = tmp_path / "cobrancas.docx"
    _make_source_docx(source)

    result = asyncio.run(resident.fill_spreadsheet_template_direct(
        str(source), str(tmp_path / "nao-existe.xlsx"),
    ))

    assert result["ok"] is False
    assert "não encontrado" in result["error"]
    resident.runtime.execute.assert_not_called()


def test_invalid_json_after_retries_fails_cleanly(tmp_path):
    resident = _make_resident()
    source = tmp_path / "cobrancas.docx"
    template = tmp_path / "modelo.xlsx"
    _make_source_docx(source)
    _make_template_xlsx(template)

    resident.runtime.execute = AsyncMock(return_value=ExecutionResult(
        plan_id="p1", status=ExecutionStatus.SUCCESS, output="isto não é JSON de jeito nenhum",
    ))

    result = asyncio.run(resident.fill_spreadsheet_template_direct(str(source), str(template)))

    assert result["ok"] is False
    assert "JSON" in result["error"]
    assert resident.runtime.execute.call_count == 2  # 1 tentativa + 1 retry


def test_mismatched_sheet_name_falls_back_to_first_sheet(tmp_path):
    resident = _make_resident()
    source = tmp_path / "cobrancas.docx"
    template = tmp_path / "modelo.xlsx"
    _make_source_docx(source)
    _make_template_xlsx(template)

    resident.runtime.execute = AsyncMock(return_value=_json_plan_response(
        "NomeErrado", [{"Nome": "Ana Silva", "CPF": "111", "Valor": "100"}],
    ))

    result = asyncio.run(resident.fill_spreadsheet_template_direct(str(source), str(template)))

    assert result["ok"] is True
    assert result["sheet_used"] == "Cobranças"


def test_runtime_failure_propagates_real_error(tmp_path):
    resident = _make_resident()
    source = tmp_path / "cobrancas.docx"
    template = tmp_path / "modelo.xlsx"
    _make_source_docx(source)
    _make_template_xlsx(template)

    resident.runtime.execute = AsyncMock(return_value=ExecutionResult(
        plan_id="p1", status=ExecutionStatus.FAILED, errors=["SpreadsheetFiller: falha real do runtime."],
    ))

    result = asyncio.run(resident.fill_spreadsheet_template_direct(str(source), str(template)))

    assert result["ok"] is False
    assert "SpreadsheetFiller" in result["error"]


def test_empty_rows_from_model_fails_cleanly(tmp_path):
    resident = _make_resident()
    source = tmp_path / "cobrancas.docx"
    template = tmp_path / "modelo.xlsx"
    _make_source_docx(source)
    _make_template_xlsx(template)

    resident.runtime.execute = AsyncMock(return_value=_json_plan_response("Cobranças", []))

    result = asyncio.run(resident.fill_spreadsheet_template_direct(str(source), str(template)))

    assert result["ok"] is False
    assert "nenhuma linha" in result["error"].lower()

"""Fase 10 - XLSX Output Writer: testes.

Cobre exatamente os cenários pedidos pelo usuário na especificação de
Fase 10, com foco especial em PRESERVAÇÃO DE TEMPLATE REAL - "Isso vale
mais para a Phoenix do que provar apenas que openpyxl.Workbook() consegue
criar um .xlsx." O arquivo `.xlsx` real do MarketUP ainda não foi
anexado pelo usuário; até lá, este arquivo usa um FIXTURE SINTÉTICO
construído via `openpyxl` com estrutura realista (múltiplas colunas,
cabeçalho formatado, larguras definidas, uma fórmula, nome de aba
específico) - exatamente a exigência de teste que o usuário descreveu.

Nenhum arquivo protegido foi tocado - este arquivo é novo, isolado,
dentro de `TESTS/`.
"""
import hashlib
import os
import tempfile

import openpyxl
import pytest
from openpyxl.styles import Font

from phoenix_kernel.documents.evidence_engine import EvidenceEntry, EvidenceStatus, FieldEvidence, ValueGroup
from phoenix_kernel.documents.identity_engine import CanonicalRecord
from phoenix_kernel.documents.normalized import ColumnMapping
from phoenix_kernel.documents.output_writer import (
    _AUDIT_SHEET_NAME,
    write_blank_workbook,
    write_to_template,
)

_SHEET_NAME = "Produtos"
_HEADERS = [
    "ID Interno",
    "Nome do Produto",
    "Preço de Venda",
    "NCM",
    "CEST",
    "Código de Barras",
    "Peso Bruto",
    "Marca",
    "Tags",
    "Estoque Calculado",
]
_HEADER_ROW = 1
_FORMULA_COLUMN_INDEX = 10  # "Estoque Calculado" - NUNCA mapeado pelo Writer.


def _file_hash(path: str) -> str:
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


@pytest.fixture()
def template_path(tmp_path):
    """Constrói um template .xlsx realista: aba nomeada, cabeçalho em
    negrito, larguras customizadas em duas colunas, e uma fórmula
    pré-existente numa coluna que o Writer nunca vai mapear."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = _SHEET_NAME

    bold_font = Font(bold=True)
    for col_index, header in enumerate(_HEADERS, start=1):
        cell = sheet.cell(row=_HEADER_ROW, column=col_index, value=header)
        cell.font = bold_font

    sheet.column_dimensions["B"].width = 40
    sheet.column_dimensions["C"].width = 15

    # linha de exemplo pré-existente com uma fórmula numa coluna não
    # mapeada pelo Writer - deve sobreviver intacta.
    sheet.cell(row=2, column=1, value=1)
    sheet.cell(row=2, column=_FORMULA_COLUMN_INDEX, value="=A2*2")

    path = str(tmp_path / "template_marketup.xlsx")
    workbook.save(path)
    return path


def _columns() -> list[ColumnMapping]:
    return [
        ColumnMapping(header="Nome do Produto", field_type="name", required=True),
        ColumnMapping(header="Preço de Venda", field_type="price", required=False),
        ColumnMapping(header="NCM", field_type="ncm", required=False),
        ColumnMapping(header="Código de Barras", field_type="ean", required=False),
        ColumnMapping(header="Marca", field_type="brand", required=False),
    ]


def _confirmed_evidence(field_type: str, record_id: str, value) -> FieldEvidence:
    entry = EvidenceEntry(
        candidate_id="c1", source="text", value=value, valid=True, confidence=0.99,
        source_record=record_id, source_method="regex",
    )
    return FieldEvidence(
        field_type=field_type, record_id=record_id, status=EvidenceStatus.CONFIRMED.value,
        value=value, final_confidence=0.99, evidence=[entry],
    )


def _conflict_evidence(field_type: str, record_id: str, value_a, value_b) -> FieldEvidence:
    entry_a = EvidenceEntry(
        candidate_id="c1", source="text", value=value_a, valid=True, confidence=0.9,
        source_record=record_id, source_method="regex",
    )
    entry_b = EvidenceEntry(
        candidate_id="c2", source="text", value=value_b, valid=True, confidence=0.9,
        source_record=record_id, source_method="regex",
    )
    return FieldEvidence(
        field_type=field_type, record_id=record_id, status=EvidenceStatus.CONFLICT.value,
        values=[
            ValueGroup(value=value_a, evidence_count=1, confidence=0.9),
            ValueGroup(value=value_b, evidence_count=1, confidence=0.9),
        ],
        evidence=[entry_a, entry_b],
    )


def _invalid_evidence(field_type: str, record_id: str, raw_value) -> FieldEvidence:
    entry = EvidenceEntry(
        candidate_id="c1", source="text", value=raw_value, valid=False, confidence=0.0,
        source_record=record_id, source_method="regex",
    )
    return FieldEvidence(
        field_type=field_type, record_id=record_id, status=EvidenceStatus.INVALID.value,
        value=None, final_confidence=0.0, evidence=[entry],
    )


def test_confirmed_and_probable_fields_write_value(template_path, tmp_path):
    record = CanonicalRecord(
        canonical_id="cr_0001",
        source_record_ids=["r1"],
        merge_method="exact_ean",
        merge_confidence=1.0,
        fields={
            "name": _confirmed_evidence("name", "r1", "Caneta Azul"),
            "price": _confirmed_evidence("price", "r1", 4.5),
            "ean": FieldEvidence(
                field_type="ean", record_id="r1", status=EvidenceStatus.PROBABLE.value,
                value="7891019125302", final_confidence=0.8,
                evidence=[EvidenceEntry(candidate_id="c3", source="text", value="7891019125302", valid=True, confidence=0.8, source_record="r1")],
            ),
        },
    )
    output_path = str(tmp_path / "saida.xlsx")
    result = write_to_template(template_path, output_path, [record], _columns(), header_row=_HEADER_ROW)

    assert result.rows_written == 1
    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[_SHEET_NAME]
    assert sheet.cell(row=2, column=2).value == "Caneta Azul"
    assert sheet.cell(row=2, column=3).value == 4.5
    assert sheet.cell(row=2, column=6).value == "7891019125302"
    # NCM e Marca não tinham evidência nenhuma e não são required -> vazio, sem auditoria.
    assert sheet.cell(row=2, column=4).value is None
    assert sheet.cell(row=2, column=8).value is None
    assert result.audit_rows == []


def test_conflict_field_leaves_cell_blank_and_writes_audit(template_path, tmp_path):
    record = CanonicalRecord(
        canonical_id="cr_0042",
        source_record_ids=["r12", "r87"],
        merge_method="exact_ean",
        merge_confidence=1.0,
        fields={
            "name": _confirmed_evidence("name", "r12", "Produto X"),
            "price": _conflict_evidence("price", "r12", 19.90, 21.90),
        },
    )
    output_path = str(tmp_path / "saida.xlsx")
    result = write_to_template(template_path, output_path, [record], _columns(), header_row=_HEADER_ROW)

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[_SHEET_NAME]
    # célula de negócio (Preço de Venda, coluna C) fica vazia - nunca um marcador de texto.
    assert sheet.cell(row=2, column=3).value is None

    assert len(result.audit_rows) == 1
    audit = result.audit_rows[0]
    assert audit.canonical_id == "cr_0042"
    assert audit.row_number == 2
    assert audit.column_header == "Preço de Venda"
    assert audit.field_type == "price"
    assert audit.status == "conflict"
    assert audit.alternatives == "19.9 | 21.9"
    assert audit.source_record_ids == "r12,r87"

    audit_sheet = workbook[_AUDIT_SHEET_NAME]
    assert audit_sheet.cell(row=1, column=1).value == "canonical_id"
    assert audit_sheet.cell(row=2, column=1).value == "cr_0042"
    assert audit_sheet.cell(row=2, column=3).value == "Preço de Venda"
    assert audit_sheet.cell(row=2, column=5).value == "conflict"


def test_invalid_field_leaves_cell_blank_and_writes_audit(template_path, tmp_path):
    record = CanonicalRecord(
        canonical_id="cr_0002",
        source_record_ids=["r5"],
        merge_method="exact_ean",
        merge_confidence=1.0,
        fields={
            "name": _confirmed_evidence("name", "r5", "Produto Y"),
            "ean": _invalid_evidence("ean", "r5", "1234567890000"),
        },
    )
    output_path = str(tmp_path / "saida.xlsx")
    result = write_to_template(template_path, output_path, [record], _columns(), header_row=_HEADER_ROW)

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[_SHEET_NAME]
    assert sheet.cell(row=2, column=6).value is None
    assert len(result.audit_rows) == 1
    assert result.audit_rows[0].status == "invalid"
    assert result.audit_rows[0].field_type == "ean"


def test_missing_required_field_leaves_blank_and_audits_without_aborting(template_path, tmp_path):
    # "name" é required=True em _columns() e não existe em fields.
    record = CanonicalRecord(
        canonical_id="cr_0003",
        source_record_ids=["r9"],
        merge_method="exact_ean",
        merge_confidence=1.0,
        fields={
            "price": _confirmed_evidence("price", "r9", 10.0),
        },
    )
    output_path = str(tmp_path / "saida.xlsx")
    result = write_to_template(template_path, output_path, [record], _columns(), header_row=_HEADER_ROW)

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[_SHEET_NAME]
    assert sheet.cell(row=2, column=2).value is None  # Nome do Produto vazio
    assert sheet.cell(row=2, column=3).value == 10.0  # price escreveu normal

    missing_audit = [a for a in result.audit_rows if a.status == "missing_required"]
    assert len(missing_audit) == 1
    assert missing_audit[0].column_header == "Nome do Produto"
    assert missing_audit[0].field_type == "name"
    # o workbook inteiro NÃO foi abortado - a linha existe, outros campos escreveram normal.


def test_missing_optional_field_is_blank_without_audit(template_path, tmp_path):
    record = CanonicalRecord(
        canonical_id="cr_0004",
        source_record_ids=["r1"],
        merge_method="exact_ean",
        merge_confidence=1.0,
        fields={
            "name": _confirmed_evidence("name", "r1", "Produto Z"),
            # "brand" (Marca, required=False) nunca aparece em fields.
        },
    )
    output_path = str(tmp_path / "saida.xlsx")
    result = write_to_template(template_path, output_path, [record], _columns(), header_row=_HEADER_ROW)

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[_SHEET_NAME]
    assert sheet.cell(row=2, column=8).value is None  # Marca
    assert result.audit_rows == []


def test_multiple_records_write_sequential_rows(template_path, tmp_path):
    records = [
        CanonicalRecord(
            canonical_id=f"cr_{i}",
            source_record_ids=[f"r{i}"],
            merge_method="exact_ean",
            merge_confidence=1.0,
            fields={"name": _confirmed_evidence("name", f"r{i}", f"Produto {i}")},
        )
        for i in range(3)
    ]
    output_path = str(tmp_path / "saida.xlsx")
    result = write_to_template(template_path, output_path, records, _columns(), header_row=_HEADER_ROW)

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[_SHEET_NAME]
    assert result.rows_written == 3
    assert sheet.cell(row=2, column=2).value == "Produto 0"
    assert sheet.cell(row=3, column=2).value == "Produto 1"
    assert sheet.cell(row=4, column=2).value == "Produto 2"


def test_refuses_to_overwrite_template_path(template_path):
    record = CanonicalRecord(
        canonical_id="cr_1", source_record_ids=["r1"], merge_method="exact_ean",
        merge_confidence=1.0, fields={"name": _confirmed_evidence("name", "r1", "X")},
    )
    with pytest.raises(ValueError):
        write_to_template(template_path, template_path, [record], _columns(), header_row=_HEADER_ROW)


def test_unknown_header_raises_explicit_error(template_path, tmp_path):
    bad_columns = [ColumnMapping(header="Coluna Que Não Existe", field_type="x")]
    record = CanonicalRecord(
        canonical_id="cr_1", source_record_ids=["r1"], merge_method="exact_ean",
        merge_confidence=1.0, fields={},
    )
    output_path = str(tmp_path / "saida.xlsx")
    with pytest.raises(ValueError, match="Coluna Que Não Existe"):
        write_to_template(template_path, output_path, [record], bad_columns, header_row=_HEADER_ROW)


def test_unsupported_conflict_policy_raises_not_implemented(template_path, tmp_path):
    record = CanonicalRecord(
        canonical_id="cr_1", source_record_ids=["r1"], merge_method="exact_ean",
        merge_confidence=1.0, fields={"name": _confirmed_evidence("name", "r1", "X")},
    )
    output_path = str(tmp_path / "saida.xlsx")
    with pytest.raises(NotImplementedError):
        write_to_template(
            template_path, output_path, [record], _columns(), header_row=_HEADER_ROW,
            conflict_policy="skip_row",
        )


def test_template_preservation_end_to_end(template_path, tmp_path):
    """O teste mais importante desta fase, pedido explicitamente pelo
    usuário: depois do Writer rodar, o TEMPLATE ORIGINAL continua
    intacto byte a byte, e o arquivo de SAÍDA preserva estrutura, nomes
    de aba, formatação, larguras e a fórmula pré-existente."""
    hash_before = _file_hash(template_path)

    records = [
        CanonicalRecord(
            canonical_id="cr_1", source_record_ids=["r1"], merge_method="exact_ean",
            merge_confidence=1.0, fields={"name": _confirmed_evidence("name", "r1", "Produto Teste")},
        ),
    ]
    output_path = str(tmp_path / "saida.xlsx")
    write_to_template(template_path, output_path, records, _columns(), header_row=_HEADER_ROW)

    # 1. o arquivo de template original nunca foi reescrito.
    hash_after = _file_hash(template_path)
    assert hash_before == hash_after

    # 2. reabrindo o template original: estrutura intacta.
    original = openpyxl.load_workbook(template_path)
    original_sheet = original[_SHEET_NAME]
    assert original.sheetnames == [_SHEET_NAME]
    assert [c.value for c in original_sheet[_HEADER_ROW]] == _HEADERS
    assert original_sheet.column_dimensions["B"].width == 40
    assert original_sheet.column_dimensions["C"].width == 15
    assert original_sheet.cell(row=1, column=2).font.bold is True
    assert original_sheet.cell(row=2, column=_FORMULA_COLUMN_INDEX).value == "=A2*2"

    # 3. o arquivo de saída preserva a mesma estrutura (colunas/ordem/aba/
    #    formatação/fórmula) e só ganhou os dados novos + a aba de auditoria.
    output = openpyxl.load_workbook(output_path)
    output_sheet = output[_SHEET_NAME]
    assert set(output.sheetnames) == {_SHEET_NAME, _AUDIT_SHEET_NAME}
    assert [c.value for c in output_sheet[_HEADER_ROW]] == _HEADERS
    assert output_sheet.column_dimensions["B"].width == 40
    assert output_sheet.column_dimensions["C"].width == 15
    assert output_sheet.cell(row=1, column=2).font.bold is True
    # fórmula pré-existente na coluna não mapeada sobrevive intacta, mesmo
    # linha 2 tendo recebido dados novos nas colunas mapeadas.
    assert output_sheet.cell(row=2, column=_FORMULA_COLUMN_INDEX).value == "=A2*2"
    assert output_sheet.cell(row=2, column=2).value == "Produto Teste"


def test_write_blank_workbook_fallback(tmp_path):
    record = CanonicalRecord(
        canonical_id="cr_1", source_record_ids=["r1"], merge_method="exact_ean",
        merge_confidence=1.0, fields={"name": _confirmed_evidence("name", "r1", "Produto Fallback")},
    )
    output_path = str(tmp_path / "fallback.xlsx")
    result = write_blank_workbook(output_path, [record], _columns(), sheet_name="Sheet1")

    assert os.path.exists(output_path)
    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook["Sheet1"]
    assert sheet.cell(row=1, column=1).value == "Nome do Produto"
    assert sheet.cell(row=2, column=1).value == "Produto Fallback"
    assert result.rows_written == 1

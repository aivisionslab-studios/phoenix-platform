"""
Testes de regressão pra auditoria 2026-08-20 (achados #2, #3 e #10 do
LEIA-ME): antes desta bateria, a Document Engine tinha DOIS problemas que
nenhum teste cobria:

1. /api/documents/read e /api/documents/edit resolviam o modelo via
   resident.registry.resolve(), mas executavam com kernel.runtime.execute()
   DIRETO - nunca passavam por um método "_direct" do ResidentManager
   (como generate_image_direct/transcribe_direct já faziam), então
   documentos não participavam do _thermal_guard nem do
   _track_model_loaded/_active_models (VRAM Guard/Hot-Swap). Isso foi
   corrigido com ResidentManager.read_document_direct()/
   edit_document_direct(), testados aqui de ponta a ponta (extração real
   de DOCX/XLSX/PPTX + runtime.execute mockado, igual ao padrão de
   test_ollama_second_option.py).
2. /api/documents/ingest era um fluxo paralelo inferior (só PDF/TXT,
   chamava kernel.runtime.execute() direto). Virou um alias fino de
   /api/documents/read - testado aqui garantindo que os dois produzem
   exatamente o mesmo resultado pro mesmo arquivo.

Rodar com: pytest -q (de dentro de 'PHOENIX 3.0/')
"""
import asyncio
import io
import sys
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.domain.execution import ExecutionResult, ExecutionStatus
from phoenix_kernel.documents.engine import extract_text, rebuild_document
from phoenix_kernel.resident.resident_manager import ResidentManager
from phoenix_kernel.models.registry import ModelRegistry
from phoenix_kernel.orchestration.execution_arbiter import default_execution_arbiter


class _FakeUploadFile:
    """Duck-type mínimo de fastapi.UploadFile - as rotas de documento só
    usam .filename e .file (via shutil.copyfileobj), então não precisamos
    de uma instância real do Starlette pra testar a lógica das rotas."""

    def __init__(self, filename: str, data: bytes):
        self.filename = filename
        self.file = io.BytesIO(data)


def _make_resident_for_document_tests() -> ResidentManager:
    """Constrói um ResidentManager real (registry real, lendo o
    catalog/models.json do projeto) mas com runtime mockado - mesmo
    padrão de test_ollama_second_option.py: não queremos subir um
    llama.cpp/llama-server de verdade só pra testar a ponte."""
    resident = object.__new__(ResidentManager)
    resident.registry = ModelRegistry()
    resident.runtime = AsyncMock()
    resident.logs = AsyncMock()
    resident.logs.add_event = lambda *a, **k: None
    resident.ahde = None
    resident._active_models = {}
    # PHX-FIX (31/08): object.__new__() pula o __init__ real, que é o
    # único lugar que seta self.execution_arbiter (ver
    # ResidentManager.__init__) - toda chamada que passa por
    # _resolve_document_decision()/_execute_document_plan_routed()
    # quebrava com AttributeError. Usa a mesma instância default que o
    # Kernel injeta em produção.
    resident.execution_arbiter = default_execution_arbiter
    return resident


# ---------------------------------------------------------------------
# 1. Document Engine (extract_text/rebuild_document) - round-trip real
#    pros formatos binários (a auditoria só tinha confirmado TXT/MD).
# ---------------------------------------------------------------------

def test_docx_extract_and_rebuild_roundtrip(tmp_path):
    from docx import Document as DocxDocument

    src = tmp_path / "origem.docx"
    doc = DocxDocument()
    doc.add_paragraph("Parágrafo de teste da Phoenix Document Engine.")
    doc.save(str(src))

    extracted = extract_text(src)
    assert "Parágrafo de teste da Phoenix Document Engine." in extracted

    out = tmp_path / "saida.docx"
    rebuild_document("Texto reescrito pela IA.", out)
    assert out.exists()
    assert "Texto reescrito pela IA." in extract_text(out)


def test_xlsx_extract_and_rebuild_roundtrip(tmp_path):
    from openpyxl import Workbook

    src = tmp_path / "origem.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "PhoenixCell"
    wb.save(str(src))

    extracted = extract_text(src)
    assert "PhoenixCell" in extracted

    out = tmp_path / "saida.xlsx"
    rebuild_document("linha1\tlinha2", out)
    assert out.exists()


def test_pptx_extract_and_rebuild_roundtrip(tmp_path):
    from pptx import Presentation

    src = tmp_path / "origem.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Slide da Phoenix"
    prs.save(str(src))

    extracted = extract_text(src)
    assert "Slide da Phoenix" in extracted

    out = tmp_path / "saida.pptx"
    rebuild_document("Conteúdo reescrito", out)
    assert out.exists()


# ---------------------------------------------------------------------
# 2. ResidentManager.read_document_direct()/edit_document_direct() -
#    a ponte que faltava (achado #3): API -> ResidentManager -> Runtime.
# ---------------------------------------------------------------------

def test_read_document_direct_uses_resolved_model_and_tracks_it(tmp_path):
    resident = _make_resident_for_document_tests()
    doc_path = tmp_path / "relatorio.txt"
    doc_path.write_text("Conteúdo do relatório para leitura.", encoding="utf-8")

    resident.runtime.execute = AsyncMock(return_value=ExecutionResult(
        plan_id="p1", status=ExecutionStatus.SUCCESS, output="Resumo gerado pela IA.",
    ))

    result = asyncio.run(resident.read_document_direct(str(doc_path), "Resuma o documento."))

    assert result["ok"] is True
    assert result["text"] == "Resumo gerado pela IA."
    assert result["extracted_length"] > 0

    # A ponte precisa ter chamado runtime.execute com o modelo/roteador
    # resolvido de verdade pelo Model Registry (role "reasoning"), não um
    # literal solto - e precisa ter rastreado o modelo como ativo (Hot-Swap
    # / _track_model_loaded), coisa que kernel.runtime.execute() direto
    # (o jeito antigo) nunca fazia.
    sent_plan = resident.runtime.execute.call_args[0][0]
    resolved = resident.registry.resolve("reasoning")
    assert sent_plan.runtime == resolved.runtime
    assert sent_plan.model == resolved.id
    assert "Conteúdo do relatório" in sent_plan.parameters["user_prompt"]
    assert resident._active_models.get(resolved.runtime) == resolved.id


def test_read_document_direct_missing_file_fails_without_calling_runtime(tmp_path):
    resident = _make_resident_for_document_tests()
    resident.runtime.execute = AsyncMock()

    result = asyncio.run(resident.read_document_direct(str(tmp_path / "nao-existe.txt"), ""))

    assert result["ok"] is False
    assert "não encontrado" in result["error"]
    resident.runtime.execute.assert_not_called()


def test_edit_document_direct_rebuilds_file_on_disk(tmp_path, monkeypatch):
    resident = _make_resident_for_document_tests()
    doc_path = tmp_path / "original.txt"
    doc_path.write_text("Texto original.", encoding="utf-8")

    resident.runtime.execute = AsyncMock(return_value=ExecutionResult(
        plan_id="p1", status=ExecutionStatus.SUCCESS, output="Texto editado pela IA.",
    ))

    # read_document_direct/edit_document_direct escrevem em "temp/documents/output"
    # relativo ao CWD do processo - isola num tmp_path pra não sujar o repo.
    monkeypatch.chdir(tmp_path)

    result = asyncio.run(resident.edit_document_direct(str(doc_path), "Reescreva formalmente."))

    assert result["ok"] is True
    out_path = Path(result["file_path"])
    assert out_path.exists()
    assert out_path.read_text(encoding="utf-8").strip() == "Texto editado pela IA."


def test_edit_document_direct_requires_instruction(tmp_path):
    resident = _make_resident_for_document_tests()
    doc_path = tmp_path / "original.txt"
    doc_path.write_text("Texto original.", encoding="utf-8")

    result = asyncio.run(resident.edit_document_direct(str(doc_path), "  "))

    assert result["ok"] is False
    assert "Instrução" in result["error"]


# ---------------------------------------------------------------------
# 3. Rotas HTTP em api_server.py - /api/documents/ingest precisa ser um
#    alias de verdade de /api/documents/read (achado #2), e as duas rotas
#    precisam chamar o Resident em vez de kernel.runtime.execute() direto
#    (achado #3).
# ---------------------------------------------------------------------

def test_ingest_route_is_a_real_alias_of_read_route(tmp_path, monkeypatch):
    import api_server

    fake_resident = object.__new__(ResidentManager)
    fake_resident.read_document_direct = AsyncMock(return_value={
        "ok": True, "text": "Análise via ponte direta.", "extracted_length": 42,
    })
    monkeypatch.setattr(api_server, "kernel", type("K", (), {"resident": fake_resident})())
    monkeypatch.chdir(tmp_path)

    async def _run():
        r1 = await api_server.read_document(file=_FakeUploadFile("doc.txt", b"conteudo"), question="")
        r2 = await api_server.ingest_document(file=_FakeUploadFile("doc.txt", b"conteudo"))
        return r1, r2

    r1, r2 = asyncio.run(_run())

    assert r1 == r2 == {"ok": True, "text": "Análise via ponte direta.", "extracted_length": 42, "file": "doc.txt"}
    # As duas rotas precisam ter passado pela MESMA ponte do Resident -
    # nunca uma lógica de extração/execução paralela e própria.
    assert fake_resident.read_document_direct.call_count == 2


def test_read_route_calls_resident_not_runtime_execute_directly(tmp_path, monkeypatch):
    import api_server

    fake_resident = object.__new__(ResidentManager)
    fake_resident.read_document_direct = AsyncMock(return_value={"ok": True, "text": "ok", "extracted_length": 1})
    monkeypatch.setattr(api_server, "kernel", type("K", (), {"resident": fake_resident})())
    monkeypatch.chdir(tmp_path)

    with patch("api_server.kernel.runtime", create=True) as fake_runtime_execute:
        result = asyncio.run(api_server.read_document(file=_FakeUploadFile("a.txt", b"x"), question="oi"))

    assert result["ok"] is True
    fake_resident.read_document_direct.assert_awaited_once()
    call_args = fake_resident.read_document_direct.call_args
    assert call_args[0][1] == "oi"  # question repassada
    fake_runtime_execute.execute.assert_not_called()


def test_read_route_surfaces_resident_error_as_422(tmp_path, monkeypatch):
    import api_server
    from fastapi import HTTPException

    fake_resident = object.__new__(ResidentManager)
    fake_resident.read_document_direct = AsyncMock(return_value={"ok": False, "error": "Formato não suportado."})
    monkeypatch.setattr(api_server, "kernel", type("K", (), {"resident": fake_resident})())
    monkeypatch.chdir(tmp_path)

    with pytest.raises(HTTPException) as exc_info:
        asyncio.run(api_server.read_document(file=_FakeUploadFile("a.txt", b"x"), question=""))

    assert exc_info.value.status_code == 422
    assert exc_info.value.detail == "Formato não suportado."


def test_edit_route_returns_base64_and_cleans_temp_output(tmp_path, monkeypatch):
    import api_server

    monkeypatch.chdir(tmp_path)
    out_dir = tmp_path / "temp" / "documents" / "output"
    out_dir.mkdir(parents=True)
    out_file = out_dir / "editado.txt"
    out_file.write_text("conteudo editado", encoding="utf-8")

    fake_resident = object.__new__(ResidentManager)
    fake_resident.edit_document_direct = AsyncMock(return_value={
        "ok": True, "file_name": "editado.txt", "file_path": str(out_file),
    })
    monkeypatch.setattr(api_server, "kernel", type("K", (), {"resident": fake_resident})())

    result = asyncio.run(api_server.edit_document(file=_FakeUploadFile("a.txt", b"x"), instruction="reescreva"))

    assert result["ok"] is True
    assert result["mime_type"] == "text/plain"
    import base64
    assert base64.b64decode(result["file_base64"]).decode("utf-8") == "conteudo editado"
    # A rota precisa limpar o arquivo de saída temporário depois de ler -
    # não deixar lixo acumulando em temp/documents/output.
    assert not out_file.exists()
